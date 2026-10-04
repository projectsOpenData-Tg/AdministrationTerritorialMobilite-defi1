"""Construit le référentiel des 39 préfectures et sa table de correspondance entre les sources (R-01).

Pivot : les 39 préfectures du Livret 02 du RGPH-5, découpage en vigueur en 2022.
Le script rapproche les noms de chaque source par une clé sans accent ni ponctuation,
applique trois règles documentées pour les limites HDX (40 unités de niveau 2 en 2021),
et s'arrête si un nom ne trouve pas sa correspondance.

Sources lues dans data/raw/ : routes classées (D5), auto-écoles (D6), panneaux de signalisation,
recensement 2022 (CSV), Livret 02 (PDF), limites administratives HDX (D8).

Usage : python scripts/referentiel_prefectures.py
"""
import csv
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

RACINE = Path(__file__).resolve().parent.parent
BRUT = RACINE / "data" / "raw"
SORTIE = RACINE / "data" / "reference" / "referentiel_prefectures.csv"

# Unités HDX dont le nom ne correspond à aucune préfecture de 2022 : préfecture cible et raison.
REGLES_HDX = {
    "Lome Commune": ("Golfe", "fusionnée avec Golfe, car Lomé Commune n’est plus une préfecture"),
    "Plaine du Mo": ("Mô", "ancien nom de Mô"),
    "Naki-Ouest": ("Kpendjal-Ouest", "rattachée à Kpendjal-Ouest par élimination (7 préfectures des Savanes dans "
                   "les deux sources) ; libellé douteux, Naki-Ouest étant un canton de Tône au recensement : "
                   "à vérifier sur la carte en 04"),
}
# Libellés de préfecture mal orthographiés dans le recensement CSV : libellé publié → préfecture.
COQUILLES_RGPH = {"TOTOAL AVE": "Avé"}
GRAND_LOME = {"Golfe", "Agoè-Nyivé"}

COLONNES = [
    "Code HDX", "Préfecture", "Région", "Zone",
    "Nom routes classées (D5)", "Nom auto-écoles (D6)", "Nom recensement CSV (D7)",
    "Nom Livret 02 (D7)", "Nom HDX (D8)", "Codes HDX regroupés", "Remarque",
]


def cle(nom):
    sans_accent = unicodedata.normalize("NFKD", nom.lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", sans_accent)


def noms_csv(fichier, colonne):
    with open(BRUT / fichier, encoding="utf-8-sig", newline="") as f:
        return {r[colonne].strip(): r for r in csv.DictReader(f) if r[colonne].strip()}


def noms_livret02():
    texte = "\n".join(page.extract_text() or "" for page in PdfReader(BRUT / "rgph5_livret02_age_milieu_prefecture.pdf").pages)
    titres = re.findall(
        r"Effectif de la population de la préfecture\s+(?:du |de la |des |de |d’|d')\s*(.+?)\s+par groupe", texte)
    return set(titres)


def main():
    erreurs = []
    livret = noms_livret02()
    if len(livret) != 39:
        sys.exit(f"Livret 02 : {len(livret)} préfectures trouvées au lieu de 39")

    routes = noms_csv("routes_classees.csv", "prefecture_nom_bdd")
    auto_ecoles = noms_csv("auto_ecoles.csv", "prefecture_nom_bdd")
    panneaux = noms_csv("equipements_panneaux_signalisation.csv", "prefecture_nom_bdd")
    with open(BRUT / "rgph_2022_population.csv", encoding="utf-8-sig", newline="") as f:
        libelles_rgph = Counter(r["découpage-administratif"].strip() for r in csv.DictReader(f))
    hdx = {
        r[0]: {"pcode": r[4], "region": r[5]}
        for r in list(load_workbook(BRUT / "limites_administratives_hdx.xlsx", read_only=True)["tgo_admin2"]
                      .iter_rows(values_only=True))[1:]
        if r[0]
    }

    # Nom de référence : l’orthographe de la collecte nationale (panneaux, 39 préfectures, accents complets).
    reference = {cle(n): n for n in panneaux}
    for nom in livret:
        if cle(nom) not in reference:
            erreurs.append(f"Livret 02 : « {nom} » absent de la collecte nationale")

    def rattacher(noms, source):
        index = {}
        for nom in noms:
            if cle(nom) in reference:
                index[cle(nom)] = nom
            else:
                erreurs.append(f"{source} : « {nom} » ne correspond à aucune des 39 préfectures")
        return index

    i_routes = rattacher(routes, "Routes classées")
    i_auto = rattacher(auto_ecoles, "Auto-écoles")
    i_livret = rattacher(livret, "Livret 02")
    # Recensement : la préfecture porte son nom seul ; les communes ajoutent un numéro (« BLITTA 3 »).
    i_rgph = {}
    for nom in libelles_rgph:
        k = cle(COQUILLES_RGPH.get(nom, nom))
        if k in reference and not re.search(r"\d", nom):
            if k in i_rgph:
                erreurs.append(f"Recensement CSV : deux libellés pour « {reference[k]} »")
            i_rgph[k] = nom

    i_hdx, regroupes, remarques = {}, {}, {}
    for nom, info in hdx.items():
        cible = REGLES_HDX[nom][0] if nom in REGLES_HDX else nom
        k = cle(cible)
        if k not in reference:
            erreurs.append(f"HDX : « {nom} » ne correspond à aucune des 39 préfectures")
            continue
        regroupes.setdefault(k, []).append(info["pcode"])
        if nom in REGLES_HDX:
            remarques.setdefault(k, []).append(f"HDX « {nom} » ({info['pcode']}) : {REGLES_HDX[nom][1]}")
        if nom not in REGLES_HDX or cible != "Golfe":
            i_hdx[k] = nom

    lignes = []
    for k, nom in reference.items():
        manque = [s for s, i in [("Livret 02", i_livret), ("recensement CSV", i_rgph), ("HDX", i_hdx)] if k not in i]
        if manque:
            erreurs.append(f"« {nom} » absente de : {', '.join(manque)}")
            continue
        codes = sorted(regroupes[k])
        principal = next(hdx[n]["pcode"] for n in hdx if n == i_hdx[k])
        region = hdx[i_hdx[k]]["region"]
        if k in i_routes and routes[i_routes[k]]["region_nom_bdd"].strip() != region:
            erreurs.append(f"« {nom} » : région {region} (HDX) ≠ {routes[i_routes[k]]['region_nom_bdd']} (routes)")
        notes = list(remarques.get(k, []))
        if k not in i_routes:
            notes.append("aucun tronçon dans les routes classées : vrai vide ou défaut de rattachement, à vérifier en 04")
        if i_rgph[k] in COQUILLES_RGPH:
            notes.append(f"libellé « {i_rgph[k]} » dans le recensement CSV (coquille)")
        if libelles_rgph[i_rgph[k]] > 1:
            notes.append(f"libellé « {i_rgph[k]} » présent {libelles_rgph[i_rgph[k]]} fois dans le recensement CSV "
                         "(préfecture et commune ou canton homonymes)")
        lignes.append({
            "Code HDX": principal,
            "Préfecture": nom,
            "Région": region,
            "Zone": "Grand Lomé" if nom in GRAND_LOME else ("Maritime hors Grand Lomé" if region == "Maritime" else region),
            "Nom routes classées (D5)": i_routes.get(k, ""),
            "Nom auto-écoles (D6)": i_auto.get(k, ""),
            "Nom recensement CSV (D7)": i_rgph[k],
            "Nom Livret 02 (D7)": i_livret[k],
            "Nom HDX (D8)": " + ".join(n for n in hdx if hdx[n]["pcode"] in codes),
            "Codes HDX regroupés": ", ".join(codes),
            "Remarque": " ; ".join(notes),
        })

    if erreurs:
        print(f"{len(erreurs)} erreur(s), référentiel non écrit :")
        for e in erreurs:
            print(f"  - {e}")
        sys.exit(1)

    lignes.sort(key=lambda l: (l["Code HDX"][:4], cle(l["Préfecture"])))
    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    with open(SORTIE, "w", newline="", encoding="utf-8") as f:
        ecrivain = csv.DictWriter(f, fieldnames=COLONNES)
        ecrivain.writeheader()
        ecrivain.writerows(lignes)

    zones = Counter(l["Zone"] for l in lignes)
    print(f"Référentiel écrit : {SORTIE} ({len(lignes)} préfectures, {len(hdx)} unités HDX rattachées)")
    print("  Par zone : " + ", ".join(f"{z} {n}" for z, n in zones.items()))
    print(f"  Routes classées : {len(i_routes)} ; auto-écoles : {len(i_auto)} ; avec remarque : "
          f"{sum(1 for l in lignes if l['Remarque'])}")


if __name__ == "__main__":
    main()
